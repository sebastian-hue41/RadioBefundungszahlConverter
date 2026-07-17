"""Tests for core/input_handler.py — file validation, loaders, bytes/path parity."""

import io
import zipfile
from pathlib import Path

import openpyxl
import pytest

from core.input_handler import (
    FileValidationError,
    MapValidationError,
    MAX_FILE_SIZE_MB,
    MAX_MAP_COLUMNS,
    StatsValidationError,
    load_map,
    load_reference_amounts,
    load_reference_data,
    load_statistics,
    validate_xlsx_bytes,
    validate_xlsx_file,
)
from .helpers import (
    inject_vba_into_xlsx,
    make_corrupt_xlsx_bytes,
    make_map_wb,
    make_path_traversal_xlsx_bytes,
    make_reference_wb,
    make_stats_wb,
    make_zip_bomb_xlsx_bytes,
    make_zip_without_workbook_xml,
    save_wb,
    wb_to_bytes,
)


# ── TestValidateXlsxFile ───────────────────────────────────────────────────────

class TestValidateXlsxFile:
    """Path-based file validation (reused from main)."""

    def test_nonexistent_file(self, tmp_path):
        with pytest.raises(FileValidationError, match="not found"):
            validate_xlsx_file(str(tmp_path / "missing.xlsx"))

    def test_wrong_extension(self, tmp_path):
        p = tmp_path / "data.csv"
        p.write_bytes(b"data")
        with pytest.raises(FileValidationError, match="Unsupported"):
            validate_xlsx_file(str(p))

    def test_empty_file(self, tmp_path):
        p = tmp_path / "empty.xlsx"
        p.write_bytes(b"")
        with pytest.raises(FileValidationError, match="empty"):
            validate_xlsx_file(str(p))

    def test_file_too_large(self, tmp_path):
        p = tmp_path / "big.xlsx"
        # Sparse file: seek past the limit and write one byte
        with open(p, "wb") as f:
            f.seek((MAX_FILE_SIZE_MB + 1) * 1024 * 1024)
            f.write(b"\x00")
        with pytest.raises(FileValidationError, match="too large"):
            validate_xlsx_file(str(p))

    def test_corrupt_zip(self, tmp_path):
        p = tmp_path / "corrupt.xlsx"
        p.write_bytes(make_corrupt_xlsx_bytes())
        with pytest.raises(FileValidationError, match="corrupt"):
            validate_xlsx_file(str(p))

    def test_missing_workbook_xml(self, tmp_path):
        p = tmp_path / "no_wb.xlsx"
        p.write_bytes(make_zip_without_workbook_xml())
        with pytest.raises(FileValidationError, match="xl/workbook.xml"):
            validate_xlsx_file(str(p))

    def test_vba_macros_rejected(self, tmp_path):
        wb = make_stats_wb([(1, 1, "01.01.2024", "MRT", "MRT")])
        clean = wb_to_bytes(wb)
        p = tmp_path / "macro.xlsx"
        p.write_bytes(inject_vba_into_xlsx(clean))
        with pytest.raises(FileValidationError, match="VBA"):
            validate_xlsx_file(str(p))

    def test_valid_xlsx(self, tmp_path):
        wb = make_stats_wb([(1, 1, "01.01.2024", "MRT", "MRT")])
        p = tmp_path / "valid.xlsx"
        wb.save(p)
        result = validate_xlsx_file(str(p))
        assert isinstance(result, Path)

    def test_path_traversal_rejected_via_file(self, tmp_path):
        p = tmp_path / "traversal.xlsx"
        p.write_bytes(make_path_traversal_xlsx_bytes())
        with pytest.raises(FileValidationError, match="suspicious path"):
            validate_xlsx_file(str(p))

    def test_zip_bomb_per_entry_rejected_via_file(self, tmp_path):
        p = tmp_path / "bomb.xlsx"
        p.write_bytes(make_zip_bomb_xlsx_bytes())
        with pytest.raises(FileValidationError, match="compression ratio"):
            validate_xlsx_file(str(p))


# ── TestValidateXlsxBytes (web-specific) ──────────────────────────────────────

class TestValidateXlsxBytes:
    """validate_xlsx_bytes — validates raw bytes without a file on disk."""

    def test_empty_bytes_rejected(self):
        with pytest.raises(FileValidationError, match="empty"):
            validate_xlsx_bytes(b"")

    def test_too_large_rejected(self):
        # Create a bytes object that exceeds the size limit
        oversized = b"\x00" * ((MAX_FILE_SIZE_MB + 1) * 1024 * 1024)
        with pytest.raises(FileValidationError, match="too large"):
            validate_xlsx_bytes(oversized)

    def test_corrupt_bytes_rejected(self):
        with pytest.raises(FileValidationError, match="corrupt"):
            validate_xlsx_bytes(make_corrupt_xlsx_bytes())

    def test_missing_workbook_xml_rejected(self):
        with pytest.raises(FileValidationError, match="xl/workbook.xml"):
            validate_xlsx_bytes(make_zip_without_workbook_xml())

    def test_vba_macros_rejected(self):
        wb = make_stats_wb([(1, 1, "01.01.2024", "MRT", "MRT")])
        data = inject_vba_into_xlsx(wb_to_bytes(wb))
        with pytest.raises(FileValidationError, match="VBA"):
            validate_xlsx_bytes(data)

    def test_path_traversal_dotdot_rejected(self):
        with pytest.raises(FileValidationError, match="suspicious path"):
            validate_xlsx_bytes(make_path_traversal_xlsx_bytes("../evil.py"))

    def test_path_traversal_absolute_rejected(self):
        with pytest.raises(FileValidationError, match="suspicious path"):
            validate_xlsx_bytes(make_path_traversal_xlsx_bytes("/etc/passwd"))

    def test_zip_bomb_per_entry_rejected(self):
        with pytest.raises(FileValidationError, match="compression ratio"):
            validate_xlsx_bytes(make_zip_bomb_xlsx_bytes())

    def test_valid_bytes_accepted(self):
        wb = make_stats_wb([(1, 1, "01.01.2024", "MRT", "MRT")])
        validate_xlsx_bytes(wb_to_bytes(wb))  # must not raise


# ── TestLoadStatistics ─────────────────────────────────────────────────────────

class TestLoadStatistics:
    """Load statistics — path-based (reused from main)."""

    def test_load_statistics_valid(self, tmp_path):
        wb = make_stats_wb([(1, 1, "01.01.2024", "MRT Schädel", "MRTS")])
        result = load_statistics(save_wb(wb, tmp_path / "s.xlsx"))
        assert result["mitarbeiter"] == "KUNZ_A"
        assert len(result["leistungen"]) == 1

    def test_load_statistics_missing_leist_kurz_column(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["E10"] = "WRONG_HEADER"
        save_wb(wb, tmp_path / "s.xlsx")
        with pytest.raises(StatsValidationError, match="Leist-Kurz"):
            load_statistics(str(tmp_path / "s.xlsx"))

    def test_load_statistics_no_data_rows(self, tmp_path):
        wb = make_stats_wb([])
        with pytest.raises(StatsValidationError, match="No valid"):
            load_statistics(save_wb(wb, tmp_path / "s.xlsx"))

    def test_load_statistics_missing_mitarbeiter(self, tmp_path):
        wb = make_stats_wb(
            [(1, 1, "01.01.2024", "MRT", "MRT")],
            mitarbeiter=None,
        )
        result = load_statistics(save_wb(wb, tmp_path / "s.xlsx"))
        assert result["mitarbeiter"] is None
        assert any("Mitarbeiter" in e for e in result["errors"])

    def test_load_statistics_missing_befunddatum(self, tmp_path):
        wb = make_stats_wb(
            [(1, 1, "01.01.2024", "MRT", "MRT")],
            befunddatum=None,
        )
        result = load_statistics(save_wb(wb, tmp_path / "s.xlsx"))
        assert result["befunddatum"] is None

    def test_load_statistics_ind2_algorithm_simple(self, tmp_path):
        """[1,2,3] → total 3."""
        rows = [
            (1, 1, "01.01.2024", "A", "A"),
            (1, 2, None,         "B", "B"),
            (1, 3, None,         "C", "C"),
        ]
        result = load_statistics(save_wb(make_stats_wb(rows), tmp_path / "s.xlsx"))
        assert result["total_leistungen"] == 3

    def test_load_statistics_ind2_algorithm_multiple_groups(self, tmp_path):
        """[1,2,3,1,2] → total 5."""
        rows = [
            (1, 1, "01.01.2024", "A", "A"),
            (1, 2, None,         "B", "B"),
            (1, 3, None,         "C", "C"),
            (2, 1, "02.01.2024", "D", "D"),
            (2, 2, None,         "E", "E"),
        ]
        result = load_statistics(save_wb(make_stats_wb(rows), tmp_path / "s.xlsx"))
        assert result["total_leistungen"] == 5

    def test_load_statistics_ind2_algorithm_single_item(self, tmp_path):
        """[1] → total 1."""
        rows = [(1, 1, "01.01.2024", "A", "A")]
        result = load_statistics(save_wb(make_stats_wb(rows), tmp_path / "s.xlsx"))
        assert result["total_leistungen"] == 1

    def test_load_statistics_ind2_algorithm_descending_boundaries(self, tmp_path):
        """[3,2,1] → total 6 (each is a group boundary)."""
        rows = [
            (1, 3, "01.01.2024", "A", "A"),
            (1, 2, None,         "B", "B"),
            (1, 1, None,         "C", "C"),
        ]
        result = load_statistics(save_wb(make_stats_wb(rows), tmp_path / "s.xlsx"))
        assert result["total_leistungen"] == 6

    def test_load_statistics_ind2_algorithm_identical_sequence(self, tmp_path):
        """[1,1,1] → total 3 (each 1 ≤ previous, so each ends a group)."""
        rows = [
            (1, 1, "01.01.2024", "A", "A"),
            (2, 1, "02.01.2024", "B", "B"),
            (3, 1, "03.01.2024", "C", "C"),
        ]
        result = load_statistics(save_wb(make_stats_wb(rows), tmp_path / "s.xlsx"))
        assert result["total_leistungen"] == 3

    def test_load_statistics_date_carryforward(self, tmp_path):
        """Rows without DokDatum inherit the date from the previous row."""
        rows = [
            (1, 1, "01.01.2024", "A", "A"),
            (1, 2, None,         "B", "B"),
        ]
        result = load_statistics(save_wb(make_stats_wb(rows), tmp_path / "s.xlsx"))
        assert all(e["year"] == 2024 for e in result["leistungen"])

    def test_load_statistics_bad_date_skipped(self, tmp_path):
        """Row with unparseable date is skipped with a warning."""
        rows = [
            (1, 1, "NOT_A_DATE", "A", "A"),
            (2, 1, "01.01.2024", "B", "B"),
        ]
        result = load_statistics(save_wb(make_stats_wb(rows), tmp_path / "s.xlsx"))
        assert len(result["leistungen"]) == 1
        assert any("cannot parse year" in e for e in result["errors"])

    def test_load_statistics_null_leist_kurz_skipped(self, tmp_path):
        """Rows with null Leist-Kurz are skipped with a warning."""
        rows = [
            (1, 1, "01.01.2024", "A", None),
            (2, 1, "02.01.2024", "B", "B"),
        ]
        result = load_statistics(save_wb(make_stats_wb(rows), tmp_path / "s.xlsx"))
        assert len(result["leistungen"]) == 1
        assert any("Leist-Kurz" in e and "empty" in e for e in result["errors"])

    def test_load_statistics_formula_cell_neutralized(self, tmp_path):
        """data_only=True means formula cells return None → row skipped, not injected."""
        wb = make_stats_wb([(1, 1, "01.01.2024", "MRT", "MRT")])
        ws = wb.active
        ws.cell(row=11, column=5).value = "=MALWARE()"
        path = save_wb(wb, tmp_path / "s.xlsx")
        with pytest.raises(StatsValidationError, match="No valid"):
            load_statistics(path)

    def test_load_statistics_count_mismatch_detection(self, tmp_path):
        """Skipped row causes len(leistungen) < total_leistungen."""
        rows = [
            (1, 1, "01.01.2024", "A", "A"),
            (1, 2, "NOT_A_DATE", "B", "B"),
        ]
        result = load_statistics(save_wb(make_stats_wb(rows), tmp_path / "s.xlsx"))
        assert len(result["leistungen"]) == 1
        assert result["total_leistungen"] == 2
        assert any("cannot parse year" in e for e in result["errors"])

    def test_load_statistics_multiple_years(self, tmp_path):
        rows = [
            (1, 1, "01.01.2023", "A", "A"),
            (2, 1, "01.01.2024", "B", "B"),
            (3, 1, "01.01.2025", "C", "C"),
        ]
        result = load_statistics(save_wb(make_stats_wb(rows), tmp_path / "s.xlsx"))
        years = {e["year"] for e in result["leistungen"]}
        assert years == {2023, 2024, 2025}


# ── TestLoadStatisticsFromBytes (web-specific) ────────────────────────────────

class TestLoadStatisticsFromBytes:
    """load_statistics accepts bytes and BytesIO in addition to file paths."""

    def test_bytes_input_accepted(self, tmp_path):
        wb = make_stats_wb([(1, 1, "01.01.2024", "MRT", "MRT")])
        result = load_statistics(wb_to_bytes(wb))
        assert len(result["leistungen"]) == 1

    def test_bytesio_input_accepted(self, tmp_path):
        wb = make_stats_wb([(1, 1, "01.01.2024", "MRT", "MRT")])
        result = load_statistics(io.BytesIO(wb_to_bytes(wb)))
        assert len(result["leistungen"]) == 1

    def test_bytes_and_path_produce_identical_results(self, tmp_path):
        rows = [
            (1, 1, "01.01.2024", "MRT Schädel", "MRTS"),
            (1, 2, None,         "CT Thorax",   "CTT"),
            (2, 1, "15.06.2024", "US Abdomen",  "USA"),
        ]
        wb = make_stats_wb(rows, mitarbeiter="DOC_X", befunddatum="01.01.2024-31.12.2024")
        path = save_wb(wb, tmp_path / "s.xlsx")
        data = wb_to_bytes(wb)

        result_path = load_statistics(path)
        result_bytes = load_statistics(data)

        assert result_path["mitarbeiter"] == result_bytes["mitarbeiter"]
        assert result_path["befunddatum"] == result_bytes["befunddatum"]
        assert len(result_path["leistungen"]) == len(result_bytes["leistungen"])
        assert result_path["total_leistungen"] == result_bytes["total_leistungen"]
        for ep, eb in zip(result_path["leistungen"], result_bytes["leistungen"]):
            assert ep["leist_kurz"] == eb["leist_kurz"]
            assert ep["year"] == eb["year"]

    def test_corrupt_bytes_raises(self):
        with pytest.raises((FileValidationError, StatsValidationError)):
            load_statistics(b"not an xlsx")

    def test_vba_bytes_rejected(self):
        wb = make_stats_wb([(1, 1, "01.01.2024", "MRT", "MRT")])
        malicious = inject_vba_into_xlsx(wb_to_bytes(wb))
        with pytest.raises(FileValidationError, match="VBA"):
            load_statistics(malicious)

    def test_path_traversal_bytes_rejected(self):
        with pytest.raises(FileValidationError, match="suspicious path"):
            load_statistics(make_path_traversal_xlsx_bytes())


# ── TestLoadMap ────────────────────────────────────────────────────────────────

class TestLoadMap:
    """Map loader — path-based (reused from main)."""

    def test_load_map_valid(self, tmp_path):
        wb = make_map_wb([("MRT", "MRT Schädel", {"DL Gefäße": 0, "MRT": 1, "CT": 0})])
        result = load_map(save_wb(wb, tmp_path / "m.xlsx"))
        assert "MRT" in result["leistungen"]
        assert result["leistungen"]["MRT"]["sections"]["MRT"] == 1

    def test_load_map_b1_wrong(self, tmp_path):
        wb = make_map_wb([], b1="WRONG")
        with pytest.raises(MapValidationError, match="Leist0"):
            load_map(save_wb(wb, tmp_path / "m.xlsx"))

    def test_load_map_f1_wrong(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=1, column=2).value = "Leist0"
        ws.cell(row=1, column=6).value = "WRONG_SECTION"
        with pytest.raises(MapValidationError, match="DL Gefäße"):
            load_map(save_wb(wb, tmp_path / "m.xlsx"))

    def test_load_map_no_sections(self, tmp_path):
        """Map with valid header but no data rows → no leistung entries error."""
        wb = make_map_wb([], sections=["DL Gefäße", "MRT", "CT"])
        with pytest.raises(MapValidationError):
            load_map(save_wb(wb, tmp_path / "m.xlsx"))

    def test_load_map_duplicate_no_values(self, tmp_path):
        """Duplicate with no section values → first entry kept."""
        wb = make_map_wb([
            ("MRT", "desc1", {"DL Gefäße": 0, "MRT": 0, "CT": 0}),
            ("MRT", "desc2", {"DL Gefäße": 0, "MRT": 0, "CT": 0}),
        ])
        result = load_map(save_wb(wb, tmp_path / "m.xlsx"))
        assert "MRT" in result["leistungen"]
        assert "MRT" not in result["merged_duplicates"]

    def test_load_map_duplicate_one_has_values(self, tmp_path):
        """Duplicate where only one has section values → that entry wins."""
        wb = make_map_wb([
            ("CODE", "desc1", {"DL Gefäße": 0, "MRT": 0, "CT": 0}),
            ("CODE", "desc2", {"DL Gefäße": 1, "MRT": 0, "CT": 0}),
        ])
        result = load_map(save_wb(wb, tmp_path / "m.xlsx"))
        assert result["leistungen"]["CODE"]["sections"]["DL Gefäße"] == 1

    def test_load_map_duplicate_conflicting_merged(self, tmp_path):
        """Duplicate with conflicting assignments → merged by max()."""
        wb = make_map_wb([
            ("CODE", "desc1", {"DL Gefäße": 1, "MRT": 0, "CT": 0}),
            ("CODE", "desc2", {"DL Gefäße": 0, "MRT": 1, "CT": 0}),
        ])
        result = load_map(save_wb(wb, tmp_path / "m.xlsx"))
        assert result["leistungen"]["CODE"]["sections"]["DL Gefäße"] == 1
        assert result["leistungen"]["CODE"]["sections"]["MRT"] == 1
        assert "CODE" in result["merged_duplicates"]

    def test_load_map_duplicate_identical(self, tmp_path):
        """Duplicate with identical assignments → first entry, no merged notice."""
        wb = make_map_wb([
            ("CODE", "d1", {"DL Gefäße": 1, "MRT": 0, "CT": 0}),
            ("CODE", "d2", {"DL Gefäße": 1, "MRT": 0, "CT": 0}),
        ])
        result = load_map(save_wb(wb, tmp_path / "m.xlsx"))
        assert "CODE" not in result["merged_duplicates"]

    def test_load_map_underscore_columns_excluded(self, tmp_path):
        wb = make_map_wb(
            [("MRT", "d", {"DL Gefäße": 1, "_TRENNER": 0})],
            sections=["DL Gefäße", "_TRENNER"],
        )
        result = load_map(save_wb(wb, tmp_path / "m.xlsx"), include_underscore_columns=False)
        assert "_TRENNER" not in result["sections"]

    def test_load_map_underscore_columns_included(self, tmp_path):
        wb = make_map_wb(
            [("MRT", "d", {"DL Gefäße": 1, "_TRENNER": 1})],
            sections=["DL Gefäße", "_TRENNER"],
        )
        result = load_map(save_wb(wb, tmp_path / "m.xlsx"), include_underscore_columns=True)
        assert "_TRENNER" in result["sections"]

    def test_load_map_formula_injection_leist0(self, tmp_path):
        wb = make_map_wb([("+MALWARE()", "desc", {"DL Gefäße": 1, "MRT": 0, "CT": 0})])
        with pytest.raises(FileValidationError, match="injection"):
            load_map(save_wb(wb, tmp_path / "m.xlsx"))


# ── TestLoadMapFromBytes (web-specific) ───────────────────────────────────────

class TestLoadMapFromBytes:
    """load_map accepts bytes and BytesIO."""

    def test_bytes_input_accepted(self):
        wb = make_map_wb([("MRT", "MRT Schädel", {"DL Gefäße": 0, "MRT": 1, "CT": 0})])
        result = load_map(wb_to_bytes(wb))
        assert "MRT" in result["leistungen"]

    def test_bytesio_input_accepted(self):
        wb = make_map_wb([("MRT", "MRT Schädel", {"DL Gefäße": 0, "MRT": 1, "CT": 0})])
        result = load_map(io.BytesIO(wb_to_bytes(wb)))
        assert "MRT" in result["leistungen"]

    def test_bytes_and_path_produce_identical_results(self, tmp_path):
        codes = [
            ("MRT",  "MRT Schädel",  {"DL Gefäße": 0, "MRT": 1, "CT": 0}),
            ("CT",   "CT Thorax",    {"DL Gefäße": 0, "MRT": 0, "CT": 1}),
            ("ANGIO","Angiographie", {"DL Gefäße": 1, "MRT": 0, "CT": 0}),
        ]
        wb = make_map_wb(codes)
        path_result = load_map(save_wb(wb, tmp_path / "m.xlsx"))
        bytes_result = load_map(wb_to_bytes(wb))

        assert path_result["sections"] == bytes_result["sections"]
        assert set(path_result["leistungen"]) == set(bytes_result["leistungen"])
        for code in path_result["leistungen"]:
            assert (path_result["leistungen"][code]["sections"]
                    == bytes_result["leistungen"][code]["sections"])

    def test_vba_bytes_rejected(self):
        wb = make_map_wb([("MRT", "d", {"DL Gefäße": 0, "MRT": 1, "CT": 0})])
        with pytest.raises(FileValidationError, match="VBA"):
            load_map(inject_vba_into_xlsx(wb_to_bytes(wb)))


# ── TestLoadMapColumnLimit (web-specific) ─────────────────────────────────────

class TestLoadMapColumnLimit:
    """MAX_MAP_COLUMNS: maps with too many section columns are rejected."""

    def test_map_at_column_limit_accepted(self, tmp_path):
        sections = ["DL Gefäße"] + [f"SEC{i}" for i in range(MAX_MAP_COLUMNS - 1)]
        wb = make_map_wb([("MRT", "d", {})], sections=sections)
        result = load_map(save_wb(wb, tmp_path / "m.xlsx"))
        assert len(result["sections"]) == MAX_MAP_COLUMNS

    def test_map_exceeds_column_limit_rejected(self, tmp_path):
        sections = ["DL Gefäße"] + [f"SEC{i}" for i in range(MAX_MAP_COLUMNS)]
        wb = make_map_wb([("MRT", "d", {})], sections=sections)
        with pytest.raises(MapValidationError, match="exceeds"):
            load_map(save_wb(wb, tmp_path / "m.xlsx"))


# ── TestLoadReferenceAmounts ───────────────────────────────────────────────────

class TestLoadReferenceAmounts:
    """Reference amounts loader — path and bytes (reused + extended)."""

    def test_load_reference_valid(self, tmp_path):
        wb = make_reference_wb({"MRT": 3000, "CT": 4000})
        result = load_reference_amounts(save_wb(wb, tmp_path / "r.xlsx"))
        assert result == {"MRT": 3000, "CT": 4000}

    def test_load_reference_nonexistent_returns_empty(self, tmp_path):
        result = load_reference_amounts(str(tmp_path / "missing.xlsx"))
        assert result == {}

    def test_load_reference_header_detection_string(self, tmp_path):
        wb = make_reference_wb({"MRT": 3000}, include_header=True)
        result = load_reference_amounts(save_wb(wb, tmp_path / "r.xlsx"))
        assert "Leistungsbereich" not in result
        assert result["MRT"] == 3000

    def test_load_reference_header_detection_none(self, tmp_path):
        wb = make_reference_wb({"MRT": 3000}, include_header=False)
        result = load_reference_amounts(save_wb(wb, tmp_path / "r.xlsx"))
        assert result["MRT"] == 3000

    def test_load_reference_none_amounts_skipped(self, tmp_path):
        wb = make_reference_wb({"MRT": 3000})
        ws = wb.active
        ws.cell(row=3, column=1).value = "CT"
        ws.cell(row=3, column=2).value = None
        result = load_reference_amounts(save_wb(wb, tmp_path / "r.xlsx"))
        assert "CT" not in result

    def test_load_reference_negative_skipped(self, tmp_path):
        wb = make_reference_wb({"MRT": -1})
        result = load_reference_amounts(save_wb(wb, tmp_path / "r.xlsx"))
        assert result == {}

    def test_load_reference_too_large_skipped(self, tmp_path):
        wb = make_reference_wb({"MRT": 2_000_000})
        result = load_reference_amounts(save_wb(wb, tmp_path / "r.xlsx"))
        assert result == {}

    def test_load_reference_zero_skipped(self, tmp_path):
        wb = make_reference_wb({"MRT": 0})
        result = load_reference_amounts(save_wb(wb, tmp_path / "r.xlsx"))
        assert result == {}

    def test_load_reference_formula_injection_name_skipped(self, tmp_path):
        wb = make_reference_wb({})
        ws = wb.active
        ws.cell(row=1, column=1).value = "=MALWARE()"
        ws.cell(row=1, column=2).value = 1000
        result = load_reference_amounts(save_wb(wb, tmp_path / "r.xlsx"))
        assert result == {}

    def test_load_reference_vba_file_rejected(self, tmp_path):
        wb = make_reference_wb({"MRT": 3000})
        p = tmp_path / "r.xlsx"
        p.write_bytes(inject_vba_into_xlsx(wb_to_bytes(wb)))
        result = load_reference_amounts(str(p))
        assert result == {}

    def test_load_reference_corrupt_xlsx_returns_empty(self, tmp_path):
        p = tmp_path / "r.xlsx"
        p.write_bytes(make_corrupt_xlsx_bytes())
        result = load_reference_amounts(str(p))
        assert result == {}

    def test_load_reference_bytes_input_accepted(self):
        wb = make_reference_wb({"MRT": 3000, "CT": 4000})
        result = load_reference_amounts(wb_to_bytes(wb))
        assert result == {"MRT": 3000, "CT": 4000}

    def test_load_reference_bytesio_input_accepted(self):
        wb = make_reference_wb({"MRT": 3000})
        result = load_reference_amounts(io.BytesIO(wb_to_bytes(wb)))
        assert result == {"MRT": 3000}

    def test_load_reference_bytes_and_path_identical(self, tmp_path):
        wb = make_reference_wb({"MRT": 3000, "CT": 4000, "US": 2000})
        path_result = load_reference_amounts(save_wb(wb, tmp_path / "r.xlsx"))
        bytes_result = load_reference_amounts(wb_to_bytes(wb))
        assert path_result == bytes_result


class TestLoadReferenceData:
    """Tests for load_reference_data — full structure with combinations and order."""

    def test_amounts_matches_load_reference_amounts(self, tmp_path):
        wb = make_reference_wb({"MRT": 3000, "CT": 4000})
        path = save_wb(wb, tmp_path / "ref.xlsx")
        data = load_reference_data(path)
        assert data["amounts"] == load_reference_amounts(path)

    def test_order_follows_file_row_order(self, tmp_path):
        order = ["CT", "US", "MRT"]
        wb = make_reference_wb({"CT": 4000, "US": 3000, "MRT": 3000}, order=order)
        path = save_wb(wb, tmp_path / "ref.xlsx")
        data = load_reference_data(path)
        assert data["order"] == order

    def test_combinations_parsed_from_col_c(self, tmp_path):
        combos = {"Combo": ["MRT Prostata", "MRT Herz"]}
        wb = make_reference_wb({"Combo": 5000}, combinations=combos)
        path = save_wb(wb, tmp_path / "ref.xlsx")
        data = load_reference_data(path)
        assert data["combinations"] == {"Combo": ["MRT Prostata", "MRT Herz"]}

    def test_trailing_semicolon_in_components_stripped(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Section"
        ws["B1"] = "Betrag"
        ws["C1"] = "Includiert"
        ws["A2"] = "Combo"
        ws["B2"] = 1000
        ws["C2"] = "A;B;C;"   # trailing semicolon
        path = save_wb(wb, tmp_path / "ref.xlsx")
        data = load_reference_data(path)
        assert data["combinations"]["Combo"] == ["A", "B", "C"]

    def test_no_combination_col_returns_empty_combinations(self, tmp_path):
        wb = make_reference_wb({"MRT": 3000})
        path = save_wb(wb, tmp_path / "ref.xlsx")
        data = load_reference_data(path)
        assert data["combinations"] == {}

    def test_section_without_amount_still_in_order(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Section"
        ws["B1"] = "Betrag"
        ws["A2"] = "NoAmount"
        ws["B2"] = None
        ws["A3"] = "HasAmount"
        ws["B3"] = 1000
        path = save_wb(wb, tmp_path / "ref.xlsx")
        data = load_reference_data(path)
        assert "NoAmount" in data["order"]
        assert "HasAmount" in data["order"]
        assert "NoAmount" not in data["amounts"]
        assert data["amounts"]["HasAmount"] == 1000

    def test_nonexistent_file_returns_empty_structure(self):
        data = load_reference_data("/nonexistent/path.xlsx")
        assert data == {"amounts": {}, "combinations": {}, "order": []}

    def test_bytes_input_accepted(self):
        combos = {"Combo": ["A", "B"]}
        wb = make_reference_wb({"Combo": 500, "A": 0}, combinations=combos)
        data = load_reference_data(wb_to_bytes(wb))
        assert data["combinations"] == {"Combo": ["A", "B"]}
