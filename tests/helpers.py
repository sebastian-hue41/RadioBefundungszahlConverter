"""Test helpers and workbook builders for the web-integration branch."""

import io
import sys
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))


# ── Workbook builders ──────────────────────────────────────────────────────────

def make_stats_wb(
    rows: list,
    mitarbeiter: str = "KUNZ_A",
    befunddatum: str = "01.01.2024-31.12.2024",
    total_l8: int | None = None,
) -> openpyxl.Workbook:
    """
    Create a statistics workbook with the given structure.

    rows: list of (ind1, ind2, dok_datum, leistung, leist_kurz) tuples
    dok_datum can be None (carry forward from previous row)
    """
    wb = openpyxl.Workbook()
    ws = wb.active

    if mitarbeiter:
        ws["A1"] = "Mitarbeiter:"
        ws["B1"] = mitarbeiter
    if befunddatum:
        ws["A2"] = "Befunddatum:"
        ws["B2"] = befunddatum

    if total_l8 is not None:
        ws.cell(row=8, column=12).value = total_l8

    ws["A10"] = "IND1"
    ws["B10"] = "IND2"
    ws["C10"] = "DokDatum"
    ws["D10"] = "Leistung"
    ws["E10"] = "Leist-Kurz"

    for i, (ind1, ind2, dok_datum, leistung, leist_kurz) in enumerate(rows, 11):
        ws.cell(row=i, column=1).value = ind1
        ws.cell(row=i, column=2).value = ind2
        ws.cell(row=i, column=3).value = dok_datum
        ws.cell(row=i, column=4).value = leistung
        ws.cell(row=i, column=5).value = leist_kurz

    return wb


def make_map_wb(
    codes: list,
    sections: list[str] | None = None,
    b1: str = "Leist0",
) -> openpyxl.Workbook:
    """
    Create a map workbook.

    codes: list of (leist0, kurztext, {section: value, ...}) tuples
    sections: list of section names (defaults to ["DL Gefäße", "MRT", "CT"])
    """
    if sections is None:
        sections = ["DL Gefäße", "MRT", "CT"]

    wb = openpyxl.Workbook()
    ws = wb.active

    ws.cell(row=1, column=2).value = b1
    for i, sec in enumerate(sections, 6):
        ws.cell(row=1, column=i).value = sec

    for row_i, (leist0, kurztext, sec_vals) in enumerate(codes, 2):
        ws.cell(row=row_i, column=2).value = leist0
        ws.cell(row=row_i, column=3).value = kurztext
        for j, sec in enumerate(sections, 6):
            ws.cell(row=row_i, column=j).value = sec_vals.get(sec, 0)

    return wb


def make_reference_wb(
    amounts: dict[str, int],
    include_header: bool = True,
    combinations: "dict | None" = None,
    order: "list | None" = None,
) -> openpyxl.Workbook:
    """Create a reference amounts workbook with optional combination definitions."""
    wb = openpyxl.Workbook()
    ws = wb.active

    row = 1
    _combos = combinations or {}
    _order  = order if order is not None else list(amounts.keys())
    for name in _combos:
        if name not in _order:
            _order = list(_order) + [name]

    if include_header:
        ws.cell(row=row, column=1).value = "Leistungsbereich"
        ws.cell(row=row, column=2).value = "Benötigt"
        if _combos:
            ws.cell(row=row, column=3).value = "Includiert"
        row += 1

    for name in _order:
        ws.cell(row=row, column=1).value = name
        ws.cell(row=row, column=2).value = amounts.get(name)
        if name in _combos:
            ws.cell(row=row, column=3).value = ";".join(_combos[name])
        row += 1

    return wb


def save_wb(wb: openpyxl.Workbook, path: Path) -> str:
    """Save workbook to path and return the path string."""
    wb.save(path)
    return str(path)


def wb_to_bytes(wb: openpyxl.Workbook) -> bytes:
    """Serialize a workbook to bytes without touching the filesystem."""
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Adversarial file builders ──────────────────────────────────────────────────

def make_corrupt_xlsx_bytes() -> bytes:
    return b"this is not a zip file but claims to be an xlsx"


def make_zip_without_workbook_xml() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("xl/not_workbook.xml", "<xml/>")
        zf.writestr("some_file.txt", "content")
    return buf.getvalue()


def inject_vba_into_xlsx(xlsx_bytes: bytes) -> bytes:
    """Take valid xlsx bytes and inject a VBA macro entry."""
    src = io.BytesIO(xlsx_bytes)
    dst = io.BytesIO()
    with zipfile.ZipFile(src, "r") as zin:
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                zout.writestr(item, zin.read(item.filename))
            zout.writestr("xl/vbaProject.bin", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1fake vba")
    return dst.getvalue()


def make_path_traversal_xlsx_bytes(traversal_entry: str = "../evil.py") -> bytes:
    """Create a zip with a path-traversal entry name."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("xl/workbook.xml", "<workbook/>")
        zf.writestr(traversal_entry, "malicious content")
    return buf.getvalue()


def make_zip_bomb_xlsx_bytes() -> bytes:
    """Create an xlsx-shaped zip with a highly compressible entry (ratio >> 50)."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/workbook.xml", "<workbook/>")
        # 100 000 null bytes compress to ~100 bytes → ratio ≈ 1000 >> 50
        zf.writestr("xl/worksheets/bomb.xml", b"\x00" * 100_000)
    return buf.getvalue()


def make_multipart_body(fields: dict[str, bytes]) -> tuple[bytes, str]:
    """Build a minimal multipart/form-data body from field_name → bytes."""
    boundary = "testboundary9876543210"
    parts = []
    for name, data in fields.items():
        header = (
            f"--{boundary}\r\n"
            f'Content-Disposition: form-data; name="{name}"; filename="{name}.xlsx"\r\n'
            f"Content-Type: application/octet-stream\r\n"
            f"\r\n"
        ).encode()
        parts.append(header + data + b"\r\n")
    body = b"".join(parts) + f"--{boundary}--\r\n".encode()
    content_type = f"multipart/form-data; boundary={boundary}"
    return body, content_type


# ── Result/data builders ───────────────────────────────────────────────────────

def make_stats_data(
    examinations: list[dict],
    errors: list[str] | None = None,
    total_leistungen: int | None = None,
    total_documents: int | None = None,
    mitarbeiter: str | None = "TEST",
    befunddatum: str | None = "2024",
) -> dict:
    return {
        "errors": errors or [],
        "leistungen": examinations,
        "total_leistungen": total_leistungen,
        "total_documents": total_documents,
        "mitarbeiter": mitarbeiter,
        "befunddatum": befunddatum,
    }


def make_map_data(
    sections: list[str],
    leistungen: dict,
    merged_duplicates: list[str] | None = None,
) -> dict:
    return {
        "sections": sections,
        "leistungen": leistungen,
        "merged_duplicates": merged_duplicates or [],
    }


def make_exam(leist_kurz: str, year: int, leistung: str | None = None) -> dict:
    return {"leist_kurz": leist_kurz, "leistung": leistung, "year": year, "ind1": 1, "row": 11}


def make_leistung_entry(sections_dict: dict[str, int]) -> dict:
    return {"kurztext": "", "sections": sections_dict, "row": 2}


def make_result(
    sections: list[str],
    year_counts: dict[str, dict[int, int]],
    years: list[int],
    errors: list[str] | None = None,
    mitarbeiter: str | None = "TEST",
    befunddatum: str | None = "2024",
    total_counted: int = 10,
    total_leistungen: int | None = 10,
    total_documents: int | None = None,
    merged_duplicates: list[str] | None = None,
) -> dict:
    data = {}
    for section in sections:
        counts = {yr: year_counts.get(section, {}).get(yr, 0) for yr in years}
        counts["Total"] = sum(counts[y] for y in years)
        data[section] = counts

    return {
        "data": data,
        "years": years,
        "errors": errors or [],
        "total_counted": total_counted,
        "total_leistungen": total_leistungen,
        "total_documents": total_documents,
        "mitarbeiter": mitarbeiter,
        "befunddatum": befunddatum,
        "merged_duplicates": merged_duplicates or [],
    }


@pytest.fixture
def tmp_xlsx(tmp_path):
    """Return a helper that saves a workbook to tmp_path and returns the path string."""
    def _save(wb, name="test.xlsx"):
        path = tmp_path / name
        wb.save(path)
        return str(path)
    return _save
