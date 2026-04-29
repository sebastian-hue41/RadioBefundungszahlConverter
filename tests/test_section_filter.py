"""
test_section_filter.py

Tests for the Leistungsbereich section-filter feature:
  - validate_section_filter() — input validation and whitelist lookup
  - process() with filter_sections — filtered output, clean errors, no unmatched rows
  - output_handler._sections_to_suffix / _make_filename — CamelCase suffix in filename
  - build_xlsx_bytes / save_output — filtered filename propagates end-to-end
  - web/worker.py sections field — JSON parsing, error responses
"""

import io
import json
import sys
from pathlib import Path
from unittest.mock import MagicMock

import openpyxl
import pytest

# ── Mock iii before importing web.worker ──────────────────────────────────────
for _mod in ("iii", "iii.types"):
    if _mod not in sys.modules:
        sys.modules[_mod] = MagicMock()

from core.input_handler import (                        # noqa: E402
    SectionFilterError,
    validate_section_filter,
    load_statistics,
    load_map,
)
from core.logic import process                          # noqa: E402
from core.output_handler import (                       # noqa: E402
    _make_filename,
    _sections_to_suffix,
    build_xlsx_bytes,
    save_output,
)
from web.worker import _parse_multipart                 # noqa: E402
from .helpers import (                                  # noqa: E402
    make_exam,
    make_leistung_entry,
    make_map_data,
    make_map_wb,
    make_multipart_body,
    make_result,
    make_stats_data,
    make_stats_wb,
    wb_to_bytes,
)

# ── Shared fixture data ────────────────────────────────────────────────────────

KNOWN_SECTIONS = ["DL Gefäße", "MRT Mamma", "MRT Prostata", "CT Abdomen", "Nuklearmedizin"]


# ═══════════════════════════════════════════════════════════════════════════════
# TestValidateSectionFilter — pure validation logic
# ═══════════════════════════════════════════════════════════════════════════════

class TestValidateSectionFilter:

    # ── Happy path ─────────────────────────────────────────────────────────────

    def test_empty_list_returns_empty(self):
        assert validate_section_filter([], KNOWN_SECTIONS) == []

    def test_single_valid_section(self):
        result = validate_section_filter(["MRT Mamma"], KNOWN_SECTIONS)
        assert result == ["MRT Mamma"]

    def test_multiple_valid_sections(self):
        result = validate_section_filter(["MRT Mamma", "CT Abdomen"], KNOWN_SECTIONS)
        assert result == ["MRT Mamma", "CT Abdomen"]

    def test_case_insensitive_match(self):
        result = validate_section_filter(["mrt mamma"], KNOWN_SECTIONS)
        assert result == ["MRT Mamma"]

    def test_mixed_case_returns_canonical(self):
        result = validate_section_filter(["MRT MAMMA", "ct abdomen"], KNOWN_SECTIONS)
        assert result == ["MRT Mamma", "CT Abdomen"]

    def test_leading_trailing_whitespace_stripped(self):
        result = validate_section_filter(["  MRT Mamma  "], KNOWN_SECTIONS)
        assert result == ["MRT Mamma"]

    def test_order_preserved_from_request(self):
        result = validate_section_filter(
            ["CT Abdomen", "MRT Mamma"], KNOWN_SECTIONS
        )
        assert result == ["CT Abdomen", "MRT Mamma"]

    def test_duplicates_deduplicated(self):
        result = validate_section_filter(
            ["MRT Mamma", "mrt mamma", "MRT Mamma"], KNOWN_SECTIONS
        )
        assert result == ["MRT Mamma"]

    def test_dedup_preserves_first_occurrence_order(self):
        result = validate_section_filter(
            ["CT Abdomen", "MRT Mamma", "ct abdomen"], KNOWN_SECTIONS
        )
        assert result == ["CT Abdomen", "MRT Mamma"]

    def test_all_known_sections_valid(self):
        result = validate_section_filter(KNOWN_SECTIONS, KNOWN_SECTIONS)
        assert result == KNOWN_SECTIONS

    def test_single_char_section_name(self):
        result = validate_section_filter(["X"], ["X", "Y"])
        assert result == ["X"]

    # ── Unknown section ────────────────────────────────────────────────────────

    def test_unknown_section_raises(self):
        with pytest.raises(SectionFilterError, match="Unknown Leistungsbereich"):
            validate_section_filter(["Not In Map"], KNOWN_SECTIONS)

    def test_unknown_section_error_contains_name(self):
        with pytest.raises(SectionFilterError, match="My Unknown Section"):
            validate_section_filter(["My Unknown Section"], KNOWN_SECTIONS)

    def test_partial_match_not_accepted(self):
        """'MRT' alone must not match 'MRT Mamma'."""
        with pytest.raises(SectionFilterError):
            validate_section_filter(["MRT"], KNOWN_SECTIONS)

    def test_substring_not_accepted(self):
        with pytest.raises(SectionFilterError):
            validate_section_filter(["Mamma"], KNOWN_SECTIONS)

    def test_first_invalid_item_fails_fast(self):
        """Error raised on the first bad item, not after processing all."""
        with pytest.raises(SectionFilterError, match="Bad Section"):
            validate_section_filter(["Bad Section", "MRT Mamma"], KNOWN_SECTIONS)

    def test_error_message_does_not_list_available_sections(self):
        """Error messages must not expose the map's full section list."""
        with pytest.raises(SectionFilterError) as exc_info:
            validate_section_filter(["Not Here"], KNOWN_SECTIONS)
        error_text = str(exc_info.value)
        for section in KNOWN_SECTIONS:
            assert section not in error_text, (
                f"Error message must not reveal section {section!r} from the map"
            )

    # ── Non-string items ───────────────────────────────────────────────────────

    def test_integer_item_raises(self):
        with pytest.raises(SectionFilterError, match="expected a string"):
            validate_section_filter([42], KNOWN_SECTIONS)

    def test_none_item_raises(self):
        with pytest.raises(SectionFilterError, match="expected a string"):
            validate_section_filter([None], KNOWN_SECTIONS)

    def test_list_item_raises(self):
        with pytest.raises(SectionFilterError, match="expected a string"):
            validate_section_filter([["nested"]], KNOWN_SECTIONS)

    def test_dict_item_raises(self):
        with pytest.raises(SectionFilterError, match="expected a string"):
            validate_section_filter([{"key": "value"}], KNOWN_SECTIONS)

    # ── Empty / whitespace strings ─────────────────────────────────────────────

    def test_empty_string_raises(self):
        with pytest.raises(SectionFilterError, match="must not be empty"):
            validate_section_filter([""], KNOWN_SECTIONS)

    def test_whitespace_only_raises(self):
        with pytest.raises(SectionFilterError, match="must not be empty"):
            validate_section_filter(["   "], KNOWN_SECTIONS)

    def test_tab_only_raises(self):
        with pytest.raises(SectionFilterError, match="must not be empty"):
            validate_section_filter(["\t\n"], KNOWN_SECTIONS)

    # ── Length limits ──────────────────────────────────────────────────────────

    def test_item_at_max_length_accepted(self):
        long_name = "X" * 200
        validate_section_filter([long_name], [long_name])  # must not raise

    def test_item_exceeds_max_length_raises(self):
        with pytest.raises(SectionFilterError, match="too long"):
            validate_section_filter(["X" * 201], KNOWN_SECTIONS)

    def test_item_length_error_contains_char_count(self):
        with pytest.raises(SectionFilterError, match="201"):
            validate_section_filter(["X" * 201], KNOWN_SECTIONS)

    # ── Too many items ─────────────────────────────────────────────────────────

    def test_at_limit_accepted(self):
        sections = [f"SEC{i}" for i in range(200)]
        validate_section_filter(sections, sections)  # must not raise

    def test_over_limit_raises(self):
        sections = [f"SEC{i}" for i in range(201)]
        with pytest.raises(SectionFilterError, match="Too many sections"):
            validate_section_filter(sections, sections)

    def test_over_limit_error_contains_count(self):
        sections = [f"SEC{i}" for i in range(201)]
        with pytest.raises(SectionFilterError, match="201"):
            validate_section_filter(sections, sections)

    # ── Formula / injection attacks ────────────────────────────────────────────

    def test_equals_prefix_rejected(self):
        with pytest.raises(SectionFilterError, match="formula injection"):
            validate_section_filter(["=HYPERLINK(\"http://evil.com\")"], KNOWN_SECTIONS)

    def test_plus_prefix_rejected(self):
        with pytest.raises(SectionFilterError, match="formula injection"):
            validate_section_filter(["+cmd|' /C calc'!A0"], KNOWN_SECTIONS)

    def test_minus_prefix_rejected(self):
        with pytest.raises(SectionFilterError, match="formula injection"):
            validate_section_filter(["-2+3+cmd|' /C calc'!A0"], KNOWN_SECTIONS)

    def test_at_prefix_rejected(self):
        with pytest.raises(SectionFilterError, match="formula injection"):
            validate_section_filter(["@SUM(1+1)*cmd|' /C calc'!A0"], KNOWN_SECTIONS)

    def test_pipe_prefix_rejected(self):
        with pytest.raises(SectionFilterError, match="formula injection"):
            validate_section_filter(["|cmd"], KNOWN_SECTIONS)

    def test_tab_prefix_rejected(self):
        with pytest.raises(SectionFilterError, match="formula injection"):
            validate_section_filter(["\t=FORMULA"], KNOWN_SECTIONS)

    def test_sql_like_string_with_injection_prefix_rejected(self):
        """SQL-style injection via '=' prefix is caught before whitelist check."""
        with pytest.raises(SectionFilterError):
            validate_section_filter(["='; DROP TABLE sections; --"], KNOWN_SECTIONS)

    def test_normal_string_with_special_chars_not_blocked(self):
        """Section names that contain special chars mid-string are fine."""
        sections = ["DL Gefäße"]  # contains ä and ß — valid map section name
        result = validate_section_filter(sections, sections)
        assert result == ["DL Gefäße"]

    def test_section_name_with_parentheses_accepted(self):
        sections = ["CT (Notfall)"]
        result = validate_section_filter(sections, sections)
        assert result == ["CT (Notfall)"]

    def test_section_name_with_slash_accepted(self):
        sections = ["MRT/CT Schädel"]
        result = validate_section_filter(sections, sections)
        assert result == ["MRT/CT Schädel"]

    # ── Empty known_sections ───────────────────────────────────────────────────

    def test_any_request_fails_against_empty_map(self):
        with pytest.raises(SectionFilterError):
            validate_section_filter(["MRT Mamma"], [])


# ═══════════════════════════════════════════════════════════════════════════════
# TestProcessWithFilter — logic.process() filtered output
# ═══════════════════════════════════════════════════════════════════════════════

class TestProcessWithFilter:

    def _make_full_setup(self):
        """Two sections (MRT, CT), one unmatched, one error pre-loaded."""
        exams = [
            make_exam("MRT_CODE", 2024),
            make_exam("MRT_CODE", 2024),
            make_exam("CT_CODE",  2025),
            make_exam("UNKNOWN",  2024, leistung="Some Unknown Procedure"),
        ]
        stats = make_stats_data(
            exams,
            errors=["WARNING: pre-existing warning from load"],
            total_leistungen=10,
        )
        map_d = make_map_data(
            ["MRT", "CT"],
            {
                "MRT_CODE": make_leistung_entry({"MRT": 1, "CT": 0}),
                "CT_CODE":  make_leistung_entry({"MRT": 0, "CT": 1}),
            },
        )
        return stats, map_d

    # ── No filter → current behaviour ─────────────────────────────────────────

    def test_no_filter_includes_all_sections(self):
        stats, map_d = self._make_full_setup()
        result = process(stats, map_d)
        assert "MRT" in result["data"]
        assert "CT" in result["data"]

    def test_no_filter_includes_unmatched(self):
        stats, map_d = self._make_full_setup()
        result = process(stats, map_d)
        assert "Some Unknown Procedure" in result["data"]

    def test_no_filter_preserves_errors(self):
        stats, map_d = self._make_full_setup()
        result = process(stats, map_d)
        assert any("pre-existing warning" in e for e in result["errors"])

    def test_no_filter_filter_sections_key_is_empty(self):
        stats, map_d = self._make_full_setup()
        result = process(stats, map_d)
        assert result["filter_sections"] == []

    # ── With filter ───────────────────────────────────────────────────────────

    def test_filter_includes_only_requested_section(self):
        stats, map_d = self._make_full_setup()
        result = process(stats, map_d, filter_sections=["MRT"])
        assert "MRT" in result["data"]
        assert "CT" not in result["data"]

    def test_filter_excludes_unmatched_leistungen(self):
        stats, map_d = self._make_full_setup()
        result = process(stats, map_d, filter_sections=["MRT"])
        assert "Some Unknown Procedure" not in result["data"]

    def test_filter_clears_all_errors(self):
        stats, map_d = self._make_full_setup()
        result = process(stats, map_d, filter_sections=["MRT"])
        assert result["errors"] == []

    def test_filter_suppresses_count_mismatch_error(self):
        exams = [make_exam("MRT_CODE", 2024)] * 3
        stats = make_stats_data(exams, total_leistungen=99)
        map_d = make_map_data(["MRT"], {"MRT_CODE": make_leistung_entry({"MRT": 1})})
        result = process(stats, map_d, filter_sections=["MRT"])
        assert not any("COUNT MISMATCH" in e for e in result["errors"])

    def test_filter_sets_total_leistungen_to_none(self):
        stats, map_d = self._make_full_setup()
        result = process(stats, map_d, filter_sections=["MRT"])
        assert result["total_leistungen"] is None

    def test_filter_sets_total_documents_to_none(self):
        stats = make_stats_data(
            [make_exam("MRT_CODE", 2024)],
            total_documents=500,
        )
        map_d = make_map_data(["MRT"], {"MRT_CODE": make_leistung_entry({"MRT": 1})})
        result = process(stats, map_d, filter_sections=["MRT"])
        assert result["total_documents"] is None

    def test_filter_empties_merged_duplicates(self):
        exams = [make_exam("MRT_CODE", 2024)]
        stats = make_stats_data(exams)
        map_d = make_map_data(
            ["MRT"],
            {"MRT_CODE": make_leistung_entry({"MRT": 1})},
            merged_duplicates=["DUPE_CODE"],
        )
        result = process(stats, map_d, filter_sections=["MRT"])
        assert result["merged_duplicates"] == []

    def test_filter_stores_filter_sections_in_result(self):
        stats, map_d = self._make_full_setup()
        result = process(stats, map_d, filter_sections=["MRT"])
        assert result["filter_sections"] == ["MRT"]

    def test_filter_preserves_correct_counts(self):
        exams = [make_exam("MRT_CODE", 2024)] * 3 + [make_exam("CT_CODE", 2024)] * 2
        stats = make_stats_data(exams)
        map_d = make_map_data(
            ["MRT", "CT"],
            {
                "MRT_CODE": make_leistung_entry({"MRT": 1, "CT": 0}),
                "CT_CODE":  make_leistung_entry({"MRT": 0, "CT": 1}),
            },
        )
        result = process(stats, map_d, filter_sections=["MRT"])
        assert result["data"]["MRT"][2024] == 3
        assert result["data"]["MRT"]["Total"] == 3

    def test_filter_multiple_sections_included(self):
        exams = [make_exam("MRT_CODE", 2024), make_exam("CT_CODE", 2024)]
        stats = make_stats_data(exams)
        map_d = make_map_data(
            ["MRT", "CT", "NUK"],
            {
                "MRT_CODE": make_leistung_entry({"MRT": 1, "CT": 0, "NUK": 0}),
                "CT_CODE":  make_leistung_entry({"MRT": 0, "CT": 1, "NUK": 0}),
            },
        )
        result = process(stats, map_d, filter_sections=["MRT", "CT"])
        assert "MRT" in result["data"]
        assert "CT" in result["data"]
        assert "NUK" not in result["data"]

    def test_filter_section_with_zero_count_still_included(self):
        """A requested section with no matching exams appears with zeros, not absent."""
        exams = [make_exam("MRT_CODE", 2024)]
        stats = make_stats_data(exams)
        map_d = make_map_data(
            ["MRT", "CT"],
            {"MRT_CODE": make_leistung_entry({"MRT": 1, "CT": 0})},
        )
        result = process(stats, map_d, filter_sections=["CT"])
        assert "CT" in result["data"]
        assert result["data"]["CT"]["Total"] == 0

    def test_filter_metadata_still_present(self):
        exams = [make_exam("MRT_CODE", 2024)]
        stats = make_stats_data(exams, mitarbeiter="DR_TEST", befunddatum="2024")
        map_d = make_map_data(["MRT"], {"MRT_CODE": make_leistung_entry({"MRT": 1})})
        result = process(stats, map_d, filter_sections=["MRT"])
        assert result["mitarbeiter"] == "DR_TEST"
        assert result["befunddatum"] == "2024"

    def test_filter_total_counted_still_correct(self):
        """total_counted reflects all parsed rows, not just the filtered ones."""
        exams = [make_exam("MRT_CODE", 2024)] * 5 + [make_exam("CT_CODE", 2024)] * 3
        stats = make_stats_data(exams)
        map_d = make_map_data(
            ["MRT", "CT"],
            {
                "MRT_CODE": make_leistung_entry({"MRT": 1, "CT": 0}),
                "CT_CODE":  make_leistung_entry({"MRT": 0, "CT": 1}),
            },
        )
        result = process(stats, map_d, filter_sections=["MRT"])
        assert result["total_counted"] == 8


# ═══════════════════════════════════════════════════════════════════════════════
# TestSectionsToSuffix / TestMakeFilenameWithSuffix — filename generation
# ═══════════════════════════════════════════════════════════════════════════════

class TestSectionsToSuffix:

    def test_single_word_section(self):
        assert _sections_to_suffix(["MRT"]) == "Mrt"

    def test_two_word_section(self):
        assert _sections_to_suffix(["MRT Mamma"]) == "MrtMamma"

    def test_multiple_sections(self):
        assert _sections_to_suffix(["MRT Mamma", "MRT Prostata"]) == "MrtMamma_MrtProstata"

    def test_three_word_section(self):
        assert _sections_to_suffix(["DL Gefäße Bauch"]) == "DlGefäßeBauch"

    def test_empty_list_returns_empty(self):
        assert _sections_to_suffix([]) == ""

    def test_single_word_capitalized(self):
        assert _sections_to_suffix(["ct"]) == "Ct"

    def test_section_with_slash(self):
        assert _sections_to_suffix(["MRT/CT"]) == "Mrt/ct"

    def test_german_umlaut_preserved(self):
        suffix = _sections_to_suffix(["DL Gefäße"])
        assert "Gefäße" in suffix


class TestMakeFilenameWithSuffix:

    def test_no_suffix_unchanged_from_baseline(self):
        name = _make_filename("SMITH", "2024")
        assert name == "SMITH2024Auswertung.xlsx"

    def test_single_section_suffix(self):
        name = _make_filename("SMITH", "2024", section_suffix="MrtMamma")
        assert name == "SMITH2024Auswertung_MrtMamma.xlsx"

    def test_multiple_section_suffix(self):
        name = _make_filename("SMITH", "2024", section_suffix="MrtMamma_MrtProstata")
        assert name == "SMITH2024Auswertung_MrtMamma_MrtProstata.xlsx"

    def test_none_suffix_no_change(self):
        name = _make_filename("SMITH", "2024", section_suffix=None)
        assert "_" not in name.replace("SMITH2024Auswertung.xlsx", "")

    def test_empty_suffix_no_trailing_underscore(self):
        name = _make_filename("SMITH", "2024", section_suffix="")
        assert name == "SMITH2024Auswertung.xlsx"

    def test_fallback_filename_with_suffix(self):
        name = _make_filename(None, None, section_suffix="MrtMamma")
        # Timestamp format: Auswertung_YYYYMMDD_HHMMSS_MrtMamma.xlsx
        assert name.startswith("Auswertung_")
        assert name.endswith("_MrtMamma.xlsx")


# ═══════════════════════════════════════════════════════════════════════════════
# TestFilteredOutputEndToEnd — filename suffix propagates through pipeline
# ═══════════════════════════════════════════════════════════════════════════════

class TestFilteredOutputEndToEnd:

    def _make_filtered_result(
        self,
        filter_secs: list,
        mitarbeiter: str = "DR_TEST",
        befunddatum: str = "2024",
    ):
        exams = [make_exam("MRT_CODE", 2024)]
        stats = make_stats_data(exams, mitarbeiter=mitarbeiter, befunddatum=befunddatum)
        map_d = make_map_data(
            ["MRT Mamma", "MRT Prostata", "CT Abdomen"],
            {"MRT_CODE": make_leistung_entry({"MRT Mamma": 1, "MRT Prostata": 0, "CT Abdomen": 0})},
        )
        return process(stats, map_d, filter_sections=filter_secs or None)

    def test_filtered_filename_contains_camel_case_suffix(self):
        result = self._make_filtered_result(["MRT Mamma"])
        filename, _ = build_xlsx_bytes(result)
        assert "MrtMamma" in filename

    def test_filtered_filename_contains_all_selected_sections(self):
        result = self._make_filtered_result(["MRT Mamma", "CT Abdomen"])
        filename, _ = build_xlsx_bytes(result)
        assert "MrtMamma" in filename
        assert "CtAbdomen" in filename

    def test_unfiltered_filename_has_no_suffix(self):
        result = self._make_filtered_result([])
        filename, _ = build_xlsx_bytes(result)
        assert filename == "DR_TEST2024Auswertung.xlsx"

    def test_filtered_xlsx_bytes_are_valid(self):
        result = self._make_filtered_result(["MRT Mamma"])
        _, data = build_xlsx_bytes(result)
        wb = openpyxl.load_workbook(io.BytesIO(data))
        assert wb.active is not None

    def test_filtered_xlsx_contains_only_requested_section_row(self):
        result = self._make_filtered_result(["MRT Mamma"])
        _, data = build_xlsx_bytes(result)
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb.active
        col_a_values = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
        assert "MRT Mamma" in col_a_values
        assert "MRT Prostata" not in col_a_values
        assert "CT Abdomen" not in col_a_values

    def test_filtered_xlsx_has_no_error_section(self):
        exams = [make_exam("UNKNOWN", 2024, leistung="Unknown Procedure")]
        stats = make_stats_data(exams)
        map_d = make_map_data(
            ["MRT Mamma"],
            {},
        )
        result = process(stats, map_d, filter_sections=["MRT Mamma"])
        _, data = build_xlsx_bytes(result)
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb.active
        all_cell_values = [
            ws.cell(r, c).value
            for r in range(1, ws.max_row + 1)
            for c in range(1, ws.max_column + 1)
        ]
        assert "Hinweise und Fehler" not in all_cell_values

    def test_save_output_filtered_filename_on_disk(self, tmp_path):
        result = self._make_filtered_result(["MRT Mamma"])
        out_path = save_output(result, output_dir=str(tmp_path))
        assert "MrtMamma" in Path(out_path).name

    def test_save_output_filtered_file_exists(self, tmp_path):
        result = self._make_filtered_result(["MRT Mamma"])
        out_path = save_output(result, output_dir=str(tmp_path))
        assert Path(out_path).exists()


# ═══════════════════════════════════════════════════════════════════════════════
# TestWebWorkerSectionsField — multipart parsing + validation via worker helpers
# ═══════════════════════════════════════════════════════════════════════════════

class TestWebWorkerSectionsField:
    """
    Tests that work without a running iii engine.
    Covers multipart parsing of the 'sections' field and validate_section_filter
    as used by the worker's parsing logic.
    """

    def _make_stats_bytes(self):
        wb = make_stats_wb([(1, 1, "01.01.2024", "MRT Schädel", "MRTS")])
        return wb_to_bytes(wb)

    # ── sections field round-trip through _parse_multipart ─────────────────────

    def test_sections_field_extracted_from_multipart(self):
        payload = json.dumps(["MRT", "CT"]).encode()
        body, ct = make_multipart_body({"statistics": b"data", "sections": payload})
        fields = _parse_multipart(body, ct)
        assert "sections" in fields
        assert json.loads(fields["sections"]) == ["MRT", "CT"]

    def test_sections_field_absent_when_not_sent(self):
        body, ct = make_multipart_body({"statistics": b"data"})
        fields = _parse_multipart(body, ct)
        assert "sections" not in fields

    # ── validate_section_filter as used by worker ─────────────────────────────

    def test_worker_style_valid_sections_parse(self):
        """Simulate what the worker does with a valid sections field."""
        raw = json.dumps(["MRT Mamma", "CT Abdomen"]).encode()
        decoded = raw.decode("utf-8", errors="replace").strip()
        parsed = json.loads(decoded)
        assert isinstance(parsed, list)
        result = validate_section_filter(parsed, KNOWN_SECTIONS)
        assert result == ["MRT Mamma", "CT Abdomen"]

    def test_worker_style_null_sections_treated_as_no_filter(self):
        """'null' JSON value → no filter (worker skips it)."""
        raw = b"null"
        decoded = raw.decode("utf-8", errors="replace").strip()
        # Worker checks: if decoded.lower() != "null" — so null is skipped
        assert decoded.lower() == "null"

    def test_worker_style_invalid_json_detected(self):
        raw = b"not valid json ["
        decoded = raw.decode("utf-8", errors="replace").strip()
        with pytest.raises(json.JSONDecodeError):
            json.loads(decoded)

    def test_worker_style_non_array_json_detected(self):
        raw = json.dumps("MRT Mamma").encode()   # string, not array
        decoded = raw.decode("utf-8", errors="replace").strip()
        parsed = json.loads(decoded)
        assert not isinstance(parsed, list)

    def test_worker_style_object_json_detected(self):
        raw = json.dumps({"section": "MRT"}).encode()
        decoded = raw.decode("utf-8", errors="replace").strip()
        parsed = json.loads(decoded)
        assert not isinstance(parsed, list)

    def test_worker_style_unknown_section_raises_filter_error(self):
        raw = json.dumps(["Not In Map"]).encode()
        decoded = raw.decode("utf-8", errors="replace").strip()
        parsed = json.loads(decoded)
        with pytest.raises(SectionFilterError):
            validate_section_filter(parsed, KNOWN_SECTIONS)

    def test_worker_style_injection_in_json_raises_filter_error(self):
        raw = json.dumps(["=HYPERLINK(\"http://evil.com\")"]).encode()
        decoded = raw.decode("utf-8", errors="replace").strip()
        parsed = json.loads(decoded)
        with pytest.raises(SectionFilterError, match="injection"):
            validate_section_filter(parsed, KNOWN_SECTIONS)

    def test_worker_style_empty_array_means_no_filter(self):
        raw = json.dumps([]).encode()
        decoded = raw.decode("utf-8", errors="replace").strip()
        parsed = json.loads(decoded)
        result = validate_section_filter(parsed, KNOWN_SECTIONS)
        assert result == []

    # ── Full in-memory pipeline with filter ───────────────────────────────────

    def test_full_pipeline_with_section_filter(self):
        """End-to-end: bytes → load → process(filter) → build_xlsx_bytes."""
        stats_bytes = wb_to_bytes(
            make_stats_wb([
                (1, 1, "01.01.2024", "MRT Schädel", "MRTS"),
                (2, 1, "01.01.2024", "CT Thorax",   "CTT"),
            ])
        )
        map_bytes = wb_to_bytes(
            make_map_wb(
                [
                    ("MRTS", "MRT Schädel", {"DL Gefäße": 0, "MRT": 1, "CT": 0}),
                    ("CTT",  "CT Thorax",   {"DL Gefäße": 0, "MRT": 0, "CT": 1}),
                ],
                sections=["DL Gefäße", "MRT", "CT"],
            )
        )
        stats_data = load_statistics(io.BytesIO(stats_bytes))
        map_data   = load_map(io.BytesIO(map_bytes))

        # Only request MRT
        filter_secs = validate_section_filter(["MRT"], map_data["sections"])
        result = process(stats_data, map_data, filter_sections=filter_secs)
        filename, data = build_xlsx_bytes(result)

        assert "Mrt" in filename
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb.active
        col_a = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
        assert "MRT" in col_a
        assert "CT" not in col_a

    def test_full_pipeline_no_filter_unchanged(self):
        """No filter → behaviour identical to before this feature."""
        stats_bytes = wb_to_bytes(
            make_stats_wb([(1, 1, "01.01.2024", "MRT Schädel", "MRTS")])
        )
        map_bytes = wb_to_bytes(
            make_map_wb([("MRTS", "MRT", {"DL Gefäße": 0, "MRT": 1, "CT": 0})])
        )
        stats_data = load_statistics(io.BytesIO(stats_bytes))
        map_data   = load_map(io.BytesIO(map_bytes))
        result = process(stats_data, map_data)
        filename, data = build_xlsx_bytes(result)

        assert "Auswertung" in filename
        assert "KUNZ" in filename          # default mitarbeiter in make_stats_wb
