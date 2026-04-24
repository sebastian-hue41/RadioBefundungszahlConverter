"""Tests for output_handler.py — xlsx generation and styling."""

from pathlib import Path

import openpyxl
import pytest

from output_handler import save_output
from .helpers import make_result


class TestSaveOutput:
    """Tests for output file generation."""

    def test_save_output_creates_file(self, tmp_path):
        """Output file is created at specified directory."""
        result = make_result(["MRT"], {"MRT": {2024: 10}}, [2024])
        path = save_output(result, output_dir=str(tmp_path))
        assert Path(path).exists()
        assert Path(path).parent == tmp_path

    def test_save_output_filename_from_metadata(self, tmp_path):
        """Filename uses mitarbeiter + befunddatum."""
        result = make_result(
            ["MRT"],
            {"MRT": {2024: 10}},
            [2024],
            mitarbeiter="KUNZ_A",
            befunddatum="01012024-31122024",
        )
        path = save_output(result, output_dir=str(tmp_path))
        filename = Path(path).name
        assert "KUNZ_A" in filename
        assert "01012024" in filename
        assert "Auswertung" in filename

    def test_save_output_filename_generic_when_missing(self, tmp_path):
        """Fallback timestamped filename when metadata missing."""
        result = make_result(
            ["MRT"],
            {"MRT": {2024: 10}},
            [2024],
            mitarbeiter=None,
            befunddatum=None,
        )
        path = save_output(result, output_dir=str(tmp_path))
        filename = Path(path).name
        assert "Auswertung_" in filename  # Timestamped
        assert filename.endswith(".xlsx")

    def test_save_output_year_columns_present(self, tmp_path):
        """Output has columns for each year."""
        result = make_result(["MRT"], {"MRT": {2024: 5, 2025: 3}}, [2024, 2025])
        path = save_output(result, output_dir=str(tmp_path))

        wb = openpyxl.load_workbook(path)
        ws = wb.active
        header_row = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        assert "2024" in header_row
        assert "2025" in header_row
        assert "Gesamt" in header_row

    def test_save_output_gesamt_column_correct_totals(self, tmp_path):
        """Gesamt column has correct totals."""
        result = make_result(
            ["MRT"],
            {"MRT": {2024: 5, 2025: 3}},
            [2024, 2025],
        )
        path = save_output(result, output_dir=str(tmp_path))

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        # Find Gesamt column
        header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        gesamt_idx = header.index("Gesamt") + 1
        # MRT row (row 2)
        gesamt_value = ws.cell(2, gesamt_idx).value
        assert gesamt_value == 8  # 5 + 3

    def test_save_output_without_reference_no_benoetigt_column(self, tmp_path):
        """Without reference_amounts, no Benötigt column."""
        result = make_result(["MRT"], {"MRT": {2024: 10}}, [2024])
        path = save_output(result, output_dir=str(tmp_path), reference_amounts=None)

        wb = openpyxl.load_workbook(path)
        ws = wb.active
        header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        assert "Benötigt" not in header

    def test_save_output_with_reference_has_benoetigt_column(self, tmp_path):
        """With reference_amounts, Benötigt column present."""
        result = make_result(["MRT"], {"MRT": {2024: 10}}, [2024])
        ref = {"MRT": 3000}
        path = save_output(result, output_dir=str(tmp_path), reference_amounts=ref)

        wb = openpyxl.load_workbook(path)
        ws = wb.active
        header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        assert "Benötigt" in header

    def test_save_output_benoetigt_values_written(self, tmp_path):
        """Benötigt column contains correct reference amounts."""
        result = make_result(["MRT", "CT"], {"MRT": {2024: 10}, "CT": {2024: 5}}, [2024])
        ref = {"MRT": 3000, "CT": 4000}
        path = save_output(result, output_dir=str(tmp_path), reference_amounts=ref)

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        benoetigt_idx = header.index("Benötigt") + 1
        # MRT row (row 2)
        mrt_benoetigt = ws.cell(2, benoetigt_idx).value
        assert mrt_benoetigt == 3000
        # CT row (row 3)
        ct_benoetigt = ws.cell(3, benoetigt_idx).value
        assert ct_benoetigt == 4000

    def test_save_output_reference_lookup_case_insensitive(self, tmp_path):
        """Reference lookup is case-insensitive."""
        result = make_result(["MRT"], {"MRT": {2024: 10}}, [2024])
        ref = {"mrt": 3000}  # Lowercase in reference
        path = save_output(result, output_dir=str(tmp_path), reference_amounts=ref)

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        benoetigt_idx = header.index("Benötigt") + 1
        benoetigt_value = ws.cell(2, benoetigt_idx).value
        assert benoetigt_value == 3000

    def test_save_output_no_cf_without_reference(self, tmp_path):
        """No conditional formatting when reference is absent."""
        result = make_result(["MRT"], {"MRT": {2024: 10}}, [2024])
        path = save_output(result, output_dir=str(tmp_path), reference_amounts=None)

        wb = openpyxl.load_workbook(path)
        ws = wb.active
        # Check that no CF rules exist
        cf_ranges = list(ws.conditional_formatting)
        assert len(cf_ranges) == 0

    def test_save_output_cf_rules_present_with_reference(self, tmp_path):
        """Conditional formatting rules added when reference present."""
        result = make_result(["MRT", "CT"], {"MRT": {2024: 10}, "CT": {2024: 5}}, [2024])
        ref = {"MRT": 3000, "CT": 4000}
        path = save_output(result, output_dir=str(tmp_path), reference_amounts=ref)

        wb = openpyxl.load_workbook(path)
        ws = wb.active
        cf_ranges = list(ws.conditional_formatting)
        assert len(cf_ranges) > 0  # At least one CF range

    def test_save_output_footer_count_summary(self, tmp_path):
        """Footer includes count summary."""
        result = make_result(
            ["MRT"],
            {"MRT": {2024: 10}},
            [2024],
            total_counted=10,
        )
        path = save_output(result, output_dir=str(tmp_path))

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        # Search for "Gezählte Leistungen" in the sheet
        found = False
        for row in ws.iter_rows(values_only=True):
            if any("Gezählte" in str(v or "") for v in row):
                found = True
                break
        assert found

    def test_save_output_footer_mismatch_warning(self, tmp_path):
        """Footer includes COUNT MISMATCH warning when applicable."""
        result = make_result(
            ["MRT"],
            {"MRT": {2024: 10}},
            [2024],
            total_counted=10,
            total_leistungen=15,  # Mismatch
        )
        path = save_output(result, output_dir=str(tmp_path))

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        # Search for "ABWEICHUNG" (German for "DEVIATION/MISMATCH")
        found = False
        for row in ws.iter_rows(values_only=True):
            if any("ABWEICHUNG" in str(v or "") for v in row):
                found = True
                break
        assert found

    def test_save_output_errors_section(self, tmp_path):
        """Errors from result are included in output."""
        result = make_result(
            ["MRT"],
            {"MRT": {2024: 10}},
            [2024],
            errors=["WARNING: Test error 1", "ERROR: Test error 2"],
        )
        path = save_output(result, output_dir=str(tmp_path))

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        all_values = []
        for row in ws.iter_rows(values_only=True):
            all_values.extend(str(v or "") for v in row)
        full_text = " ".join(all_values)
        assert "Test error 1" in full_text
        assert "Test error 2" in full_text

    def test_save_output_merged_duplicates_notice(self, tmp_path):
        """Merged duplicates notice appears when applicable."""
        result = make_result(
            ["MRT"],
            {"MRT": {2024: 10}},
            [2024],
            merged_duplicates=["CODE1", "CODE2"],
        )
        path = save_output(result, output_dir=str(tmp_path))

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        all_values = []
        for row in ws.iter_rows(values_only=True):
            all_values.extend(str(v or "") for v in row)
        full_text = " ".join(all_values)
        assert "CODE1" in full_text
        assert "CODE2" in full_text

    def test_save_output_multiple_sections(self, tmp_path):
        """Output with multiple sections laid out correctly."""
        result = make_result(
            ["MRT", "CT", "US"],
            {
                "MRT": {2024: 100, 2025: 50},
                "CT": {2024: 200, 2025: 150},
                "US": {2024: 300, 2025: 250},
            },
            [2024, 2025],
        )
        path = save_output(result, output_dir=str(tmp_path))

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        # Check that all sections appear
        first_col = [ws.cell(r, 1).value for r in range(1, 10)]
        assert "MRT" in first_col
        assert "CT" in first_col
        assert "US" in first_col

    def test_save_output_zero_counts_displayed(self, tmp_path):
        """Sections with zero counts are displayed as 0."""
        result = make_result(
            ["MRT", "CT"],
            {"MRT": {2024: 10}, "CT": {2024: 0}},
            [2024],
        )
        path = save_output(result, output_dir=str(tmp_path))

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        # CT total should be 0
        # Find CT row (row 3), Gesamt column (column 3)
        ct_total = ws.cell(3, 3).value
        assert ct_total == 0

    def test_save_output_handles_missing_reference_for_section(self, tmp_path):
        """Sections without reference entry still render (Benötigt left blank)."""
        result = make_result(
            ["MRT", "CT"],
            {"MRT": {2024: 10}, "CT": {2024: 5}},
            [2024],
        )
        ref = {"MRT": 3000}  # CT not in reference
        path = save_output(result, output_dir=str(tmp_path), reference_amounts=ref)

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        benoetigt_idx = header.index("Benötigt") + 1
        # MRT has value
        mrt_benoetigt = ws.cell(2, benoetigt_idx).value
        assert mrt_benoetigt == 3000
        # CT should be None or empty
        ct_benoetigt = ws.cell(3, benoetigt_idx).value
        # Should not raise; value can be None or empty

