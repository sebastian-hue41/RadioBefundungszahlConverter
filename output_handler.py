"""
output_handler.py — creates and saves the output xlsx file.

Designed to be server-compatible: accepts result dict + output directory,
returns the saved file path. No interactive prompts.
"""

import re
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ── Styling constants ──────────────────────────────────────────────────────────

_HEADER_BG = "2F4F8F"
_TOTAL_BG = "D9E1F2"
_ERROR_COLOR = "CC0000"
_NOTE_COLOR = "856404"
_MISMATCH_BG = "FFE0E0"

# Achievement fill colors for the "Benötigt" column (Gesamt cell coloring).
_ACH_GREEN  = "C6EFCE"   # >= 100 % of required
_ACH_YELLOW = "FFEB9C"   # >= 80 %
_ACH_ORANGE = "FFD966"   # >= 50 %
_ACH_RED    = "FFC7CE"   # <  50 %

_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _make_filename(mitarbeiter: str | None, befunddatum: str | None) -> str:
    """
    Builds the output filename from metadata.
    Falls back to a timestamped generic name when metadata is missing.
    Strips spaces; keeps alphanumeric, underscores, and hyphens only.
    """
    if mitarbeiter and befunddatum:
        clean = lambda s: re.sub(r"[^\w\-]", "", s.replace(" ", ""))
        name = f"{clean(mitarbeiter)}{clean(befunddatum)}Auswertung"
    else:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        name = f"Auswertung_{ts}"
    return f"{name}.xlsx"


def _style_header_cell(cell, value: str) -> None:
    cell.value = value
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color=_HEADER_BG, end_color=_HEADER_BG, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = _BORDER


def _style_data_cell(cell, value, bold: bool = False) -> None:
    cell.value = value
    if bold:
        cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="right" if isinstance(value, int) else "left")
    cell.border = _BORDER


def _style_total_cell(cell, value) -> None:
    cell.value = value
    cell.font = Font(bold=True)
    cell.fill = PatternFill(start_color=_TOTAL_BG, end_color=_TOTAL_BG, fill_type="solid")
    cell.alignment = Alignment(horizontal="right")
    cell.border = _BORDER


def _add_achievement_cf(ws, gesamt_col: int, benoetigt_col: int, first_row: int, last_row: int) -> None:
    """
    Add conditional-formatting rules to the Gesamt column so Excel re-colours
    cells dynamically whenever the user edits a Benötigt value.

    Rules reference the Benötigt cell in the same row via a formula, so the
    colour updates live — no need to regenerate the file after edits.
    Priority order: green (highest) → yellow → orange → red.
    """
    try:
        g = get_column_letter(gesamt_col)
        b = get_column_letter(benoetigt_col)
        cf_range = f"{g}{first_row}:{g}{last_row}"
        # Formulas are written for the first row; Excel adjusts them per row.
        gf = f"{g}{first_row}"
        bf = f"{b}{first_row}"
        for formula, color in (
            (f'AND({bf}<>"",{bf}>0,{gf}/{bf}>=1)',   _ACH_GREEN),
            (f'AND({bf}<>"",{bf}>0,{gf}/{bf}>=0.8)', _ACH_YELLOW),
            (f'AND({bf}<>"",{bf}>0,{gf}/{bf}>=0.5)', _ACH_ORANGE),
            (f'AND({bf}<>"",{bf}>0,{gf}/{bf}>0)',    _ACH_RED),
        ):
            ws.conditional_formatting.add(
                cf_range,
                FormulaRule(
                    formula=[formula],
                    fill=PatternFill(start_color=color, end_color=color, fill_type="solid"),
                    stopIfTrue=True,
                ),
            )
    except Exception:
        pass  # CF is cosmetic — never let it break the file


def _autofit_columns(ws, max_width: int = 60) -> None:
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        best = max(
            (len(str(cell.value)) for cell in col if cell.value is not None),
            default=8,
        )
        ws.column_dimensions[col_letter].width = min(best + 4, max_width)


def save_output(
    result: dict,
    output_dir: str = ".",
    reference_amounts: "dict[str, int] | None" = None,
) -> str:
    """
    Creates an xlsx file from the processing result and saves it to `output_dir`.

    Parameters
    ----------
    result      : dict returned by logic.process()
    output_dir  : directory where the file is saved (default: current directory)

    Returns
    -------
    str — absolute path of the saved file.

    Raises
    ------
    IOError if the file cannot be written.
    """
    filename = _make_filename(result.get("mitarbeiter"), result.get("befunddatum"))

    try:
        out_dir = Path(output_dir).resolve()
        out_dir.mkdir(parents=True, exist_ok=True)
        output_path = out_dir / filename
    except Exception as exc:
        raise IOError(f"Cannot prepare output directory '{output_dir}': {exc}") from exc

    years: list[int] = result["years"]
    data: dict[str, dict] = result["data"]
    errors: list[str] = result.get("errors", [])
    total_counted: int = result.get("total_counted", 0)
    total_leistungen: int | None = result.get("total_leistungen")
    total_documents: int | None = result.get("total_documents")
    merged_duplicates: list[str] = result.get("merged_duplicates", [])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Auswertung"
    ws.freeze_panes = "B2"  # freeze header row and section column

    # Build a case-insensitive fallback lookup so minor capitalisation
    # differences between the map and the reference file still match.
    _ref: dict[str, int] = reference_amounts or {}
    _ref_lower: dict[str, int] = {k.lower(): v for k, v in _ref.items()}
    has_reference = bool(_ref)

    def _lookup_required(section_name: str) -> "int | None":
        try:
            if section_name in _ref:
                return _ref[section_name]
            return _ref_lower.get(section_name.lower())
        except Exception:
            return None

    # ── Header row ────────────────────────────────────────────────────────────
    col_headers = ["Leistungsbereich"] + [str(y) for y in years] + ["Gesamt"]
    if has_reference:
        col_headers.append("Benötigt")
    for col_idx, header_text in enumerate(col_headers, 1):
        _style_header_cell(ws.cell(row=1, column=col_idx), header_text)

    ws.row_dimensions[1].height = 30

    # ── Data rows ─────────────────────────────────────────────────────────────
    gesamt_col = len(years) + 2
    benoetigt_col = gesamt_col + 1

    for row_offset, (section, year_data) in enumerate(data.items(), 2):
        _style_data_cell(ws.cell(row=row_offset, column=1), section)

        for col_offset, year in enumerate(years, 2):
            _style_data_cell(
                ws.cell(row=row_offset, column=col_offset),
                year_data.get(year, 0),
            )

        actual_total = year_data.get("Total", 0)
        gesamt_cell = ws.cell(row=row_offset, column=gesamt_col)
        _style_total_cell(gesamt_cell, actual_total)

        if has_reference:
            required = _lookup_required(section)
            try:
                if required is not None:
                    req_cell = ws.cell(row=row_offset, column=benoetigt_col)
                    req_cell.value = required
                    req_cell.font = Font(bold=True)
                    req_cell.alignment = Alignment(horizontal="right")
                    req_cell.border = _BORDER
                else:
                    ws.cell(row=row_offset, column=benoetigt_col).border = _BORDER
            except Exception:
                pass  # never let reference rendering break the output

    # Conditional formatting on the entire Gesamt column — colours update live
    # in Excel when the user edits a Benötigt cell (no file regeneration needed).
    if has_reference and data:
        _add_achievement_cf(ws, gesamt_col, benoetigt_col, first_row=2, last_row=len(data) + 1)

    next_row = len(data) + 2  # first row after data

    # ── Count summary row ─────────────────────────────────────────────────────
    next_row += 1
    ws.cell(row=next_row, column=1).value = "Gezählte Leistungen:"
    ws.cell(row=next_row, column=1).font = Font(bold=True)
    ws.cell(row=next_row, column=2).value = total_counted
    ws.cell(row=next_row, column=2).font = Font(bold=True)

    # Befunddokumente (L8) = unique IND1 count — shown for reference only.
    # It counts patient reports, not individual procedures, so no comparison is made.
    if total_documents is not None:
        next_row += 1
        ws.cell(row=next_row, column=1).value = "Befunddokumente (L8):"
        ws.cell(row=next_row, column=1).font = Font(italic=True)
        ws.cell(row=next_row, column=2).value = total_documents
        ws.cell(row=next_row, column=2).font = Font(italic=True)
        ws.cell(row=next_row, column=3).value = "(Anzahl Befundberichte / IND1 — nicht mit Leistungen vergleichbar)"
        ws.cell(row=next_row, column=3).font = Font(italic=True, color="888888")

    # Mismatch: only flag if processed rows != IND2-derived expected total.
    if total_leistungen is not None and total_counted != total_leistungen:
        skipped = total_leistungen - total_counted
        next_row += 1
        mismatch_cell = ws.cell(row=next_row, column=1)
        mismatch_cell.value = (
            f"ABWEICHUNG: IND2-Algorithmus erwartet {total_leistungen:,} Leistungen; "
            f"{skipped:,} Zeilen übersprungen"
        )
        mismatch_cell.font = Font(bold=True, color=_ERROR_COLOR)
        mismatch_cell.fill = PatternFill(
            start_color=_MISMATCH_BG, end_color=_MISMATCH_BG, fill_type="solid"
        )

    # ── Merged duplicates notice ──────────────────────────────────────────────
    if merged_duplicates:
        next_row += 2
        label_cell = ws.cell(row=next_row, column=1)
        label_cell.value = "Hinweis: Doppelte Leistungsbezeichnungen (Ergebnisse möglicherweise ungenau)"
        label_cell.font = Font(bold=True, size=11, color=_NOTE_COLOR)
        ws.merge_cells(
            start_row=next_row, start_column=1,
            end_row=next_row, end_column=len(col_headers),
        )

        next_row += 1
        desc_cell = ws.cell(row=next_row, column=1)
        desc_cell.value = (
            "Die folgenden Leistungsbezeichnungen kamen mehrfach mit widersprüchlichen "
            "Zuordnungen in der Map vor. Die Werte wurden zeilenweise per Maximum zusammengeführt. "
            "Bitte die Map-Datei prüfen und bereinigen."
        )
        desc_cell.alignment = Alignment(wrap_text=True)
        desc_cell.font = Font(color=_NOTE_COLOR)
        ws.merge_cells(
            start_row=next_row, start_column=1,
            end_row=next_row, end_column=len(col_headers),
        )
        ws.row_dimensions[next_row].height = 40

        for code in merged_duplicates:
            next_row += 1
            cell = ws.cell(row=next_row, column=1)
            cell.value = f"  - {code}"
            cell.font = Font(color=_NOTE_COLOR)

    # ── Errors / warnings section ─────────────────────────────────────────────
    if errors:
        next_row += 2
        label_cell = ws.cell(row=next_row, column=1)
        label_cell.value = "Hinweise und Fehler:"
        label_cell.font = Font(bold=True, size=11)

        for error_text in errors:
            next_row += 1
            cell = ws.cell(row=next_row, column=1)
            cell.value = error_text
            cell.alignment = Alignment(wrap_text=True)

            # Colour-code by severity prefix.
            if error_text.startswith("COUNT MISMATCH") or error_text.startswith("ERROR"):
                cell.font = Font(color=_ERROR_COLOR)
            elif error_text.startswith("WARNING"):
                cell.font = Font(color=_NOTE_COLOR)

            # Allow the error text to span multiple columns visually.
            ws.merge_cells(
                start_row=next_row,
                start_column=1,
                end_row=next_row,
                end_column=len(col_headers),
            )
            ws.row_dimensions[next_row].height = 30

    _autofit_columns(ws)

    try:
        wb.save(output_path)
    except Exception as exc:
        raise IOError(f"Failed to save output file to '{output_path}': {exc}") from exc

    return str(output_path)
